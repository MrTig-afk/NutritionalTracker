"""GET /export: the user's food log, goals, meal templates and Library as XLSX or a CSV zip.

Run:  venv/Scripts/python -m unittest backend.tests.test_export -v
No network: FakeConn scripts the DB.
"""
import csv
import io
import threading
from datetime import datetime
import os
import unittest
import zipfile
from unittest import mock

import openpyxl

import main
import export
from .test_phase0 import FakeConn, route

SHEETS = ["Food log", "Goals", "Meal templates", "Library"]
FILES = ["food-log.csv", "goals.csv", "meal-templates.csv", "library.csv"]
MILK = {"_kcal": True, "per_serving": {"size": "250 mL", "calories": 125, "protein": "9.3 g",
                                       "carbohydrates": "14.3 g", "fat": "3.3 g", "fibre": "0 g", "sodium": "105 mg"}}
VIA_CLAUDE = {**MILK, "_source": "claude", "_meal_label": "Breakfast"}
PER_100G = {"_kcal": True, "per_100g": {"calories": 400, "protein": "10 g", "sodium": "0.5 g"}}
EVIL = '=HYPERLINK("http://example.com","x")'


def full_conn():
    # the food log is read newest first (the cap keeps the newest) and written oldest first
    return FakeConn([
        ("FROM daily_log", [("2026-09-25", "Flat white", 1.0, VIA_CLAUDE),
                            ("2026-09-24", "Sungold Milk", 2.0, MILK)]),
        ("FROM user_goals", [(2200.0, 140.0, 260.0, 70.0, None)]),
        ("FROM meal_templates", [("Usual Breakfast", "Sungold Milk", 1.0, MILK), ("Empty one", None, None, None)]),
        ("FROM folders", [("Dairy", "Sungold Milk", MILK), ("Empty folder", None, None)]),
    ])


def csv_rows(body, name):
    return list(csv.reader(io.StringIO(zipfile.ZipFile(io.BytesIO(body)).read(name).decode("utf-8-sig"))))


class Export(unittest.TestCase):
    def setUp(self):
        p = mock.patch.object(export, "_recent", {})
        p.start()
        self.addCleanup(p.stop)

    def get(self, conn, fmt, **headers):
        client = route(self, conn)
        return client.get(f"/export?format={fmt}", headers={"Authorization": "Bearer x", **headers})

    def xlsx(self, conn):
        r = self.get(conn, "xlsx")
        self.assertEqual(r.status_code, 200, r.text[:200])
        self.assertEqual(r.headers["content-type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(r.headers["content-disposition"], r'^attachment; filename="nutriscan-export-\d{4}-\d{2}-\d{2}\.xlsx"$')
        self.assertEqual(r.headers["cache-control"], "no-store")
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def test_xlsx_has_the_four_sheets(self):
        wb = self.xlsx(full_conn())
        self.assertEqual(wb.sheetnames, SHEETS)
        log = list(wb["Food log"].values)
        self.assertEqual(log[0], ("Date", "Meal", "Food", "Servings", "Serving size", "Calories (kcal)", "Protein (g)",
                                  "Carbs (g)", "Fat (g)", "Fibre (g)", "Sodium (mg)", "Via Claude"))
        self.assertEqual(log[1], (datetime(2026, 9, 24), None, "Sungold Milk", 2, "250 mL", 250, 18.6, 28.6, 6.6, 0, 210, "No"))
        self.assertEqual((log[2][1], log[2][-1]), ("Breakfast", "Yes"))
        self.assertEqual(list(wb["Goals"].values)[1], (2200, 140, 260, 70, 0))   # EX-10: fibre defaults like the app
        tmpl = list(wb["Meal templates"].values)
        self.assertEqual(tmpl[1][:4], ("Usual Breakfast", "Sungold Milk", 1, 125))
        self.assertEqual(tmpl[2][:2], ("Empty one", None))   # a template with no items still shows up (empty cell)
        lib = list(wb["Library"].values)
        self.assertEqual(lib[1], ("Dairy", "Sungold Milk", "250 mL", 125, 9.3, 14.3, 3.3, 0, 105))
        self.assertEqual(lib[2][:2], ("Empty folder", None))   # EX-9: an empty folder is still there

    def test_per_100g_rows_say_so_and_sodium_is_a_total(self):
        # EX-7
        wb = self.xlsx(FakeConn([("FROM daily_log", [("2026-09-24", "Oats", 0.5, PER_100G)])]))
        row = list(wb["Food log"].values)[1]
        self.assertEqual((row[4], row[5], row[10]), ("100 g", 200, 250))

    def test_an_empty_account_still_gets_every_sheet(self):
        wb = self.xlsx(FakeConn([]))
        self.assertEqual(wb.sheetnames, SHEETS)
        self.assertEqual(len(list(wb["Food log"].values)), 1)   # header only
        self.assertEqual(list(wb["Goals"].values)[1], (2000, 150, 250, 65, 30))   # what the app shows by default

    def test_csv_zip_has_the_four_files(self):
        r = self.get(full_conn(), "csv")
        self.assertEqual(r.headers["content-type"], "application/zip")
        self.assertRegex(r.headers["content-disposition"], r'filename="nutriscan-export-\d{4}-\d{2}-\d{2}\.zip"$')
        z = zipfile.ZipFile(io.BytesIO(r.content))
        self.assertEqual(z.namelist(), FILES)
        self.assertTrue(z.read("food-log.csv").startswith(b"\xef\xbb\xbf"))   # BOM: Excel opens it as UTF-8
        rows = csv_rows(r.content, "food-log.csv")
        self.assertEqual(rows[0][:3], ["Date", "Meal", "Food"])
        self.assertEqual(rows[1][2], "Sungold Milk")

    def test_formula_like_text_stays_text(self):
        # a name is user text: it must never run as a spreadsheet formula (CSV/formula injection)
        conn = FakeConn([("FROM daily_log", [("2026-09-24", EVIL, 1.0, MILK)]), ("FROM folders", [("+cmd", "-1+1", MILK)])])
        wb = self.xlsx(conn)
        cell = wb["Food log"]["C2"]
        self.assertEqual((cell.value, cell.data_type), (EVIL, "s"))   # EX-8: a text cell, no visible apostrophe
        self.assertEqual(list(wb["Library"].values)[1][:2], ("+cmd", "-1+1"))
        rows = csv_rows(self.get(conn, "csv").content, "food-log.csv")
        self.assertEqual(rows[1][2], "'" + EVIL)   # CSV has no cell types: the apostrophe is the guard
        self.assertEqual(csv_rows(self.get(conn, "csv").content, "library.csv")[1][:2], ["'+cmd", "'-1+1"])

    def test_a_control_character_in_a_name_does_not_break_the_export(self):
        # EX-1: openpyxl refuses them outright
        conn = FakeConn([("FROM daily_log", [("2026-09-24", "milk\x0bfat\x00", 1.0, MILK)])])
        self.assertEqual(list(self.xlsx(conn)["Food log"].values)[1][2], "milk fat")

    def test_characters_xml_forbids_do_not_corrupt_the_file(self):
        # EX2-1: the file would not open at all
        conn = FakeConn([("FROM daily_log", [("2026-09-24", "x\uffffy\ufffe", 1.0, MILK)])])
        self.assertEqual(list(self.xlsx(conn)["Food log"].values)[1][2], "x y")

    def test_a_cell_never_exceeds_the_spreadsheet_limit(self):
        # EX2-8: Excel's 32,767 characters per cell
        conn = FakeConn([("FROM daily_log", [("2026-09-24", "a" * 40000, 1.0, MILK)])])
        self.assertEqual(len(list(self.xlsx(conn)["Food log"].values)[1][2]), 32767)

    def test_unknown_format(self):
        self.assertEqual(self.get(FakeConn([]), "pdf").status_code, 422)

    def test_reads_are_capped_and_scoped(self):
        conn = FakeConn([])
        self.get(conn, "xlsx")
        for sql, params in conn.executed:
            if "SELECT" in sql and "set_config" not in sql:
                self.assertEqual(params, ["user-1"])
                if "user_goals" not in sql:   # one row per user by its primary key
                    self.assertIn("LIMIT", sql)

    def test_the_cap_keeps_the_newest_and_says_so(self):
        # EX-5, EX2-7: one extra row is read to know whether anything was cut
        rows = [("2026-09-27", "d", 1.0, MILK), ("2026-09-26", "c", 1.0, MILK), ("2026-09-25", "b", 1.0, MILK)]
        conn = FakeConn([("FROM daily_log", rows)])
        with mock.patch.object(export, "LOG_CAP", 2):
            log = list(self.xlsx(conn)["Food log"].values)
        sql = [s for s, _ in conn.executed if "FROM daily_log" in s][0]
        self.assertIn("ORDER BY date DESC, created_at DESC LIMIT 3", sql)
        self.assertEqual([r[2] for r in log[1:3]], ["c", "d"])
        self.assertIn("newest 2", log[3][0])
        conn = FakeConn([("FROM daily_log", rows[:2])])   # exactly the cap: nothing was cut, no note
        with mock.patch.object(export, "LOG_CAP", 2):
            self.assertEqual(len(list(self.xlsx(conn)["Food log"].values)), 3)

    def test_templates_and_library_say_when_they_are_cut(self):
        # EX2-3
        conn = FakeConn([("FROM meal_templates", [("A", "x", 1.0, MILK)] * 3), ("FROM folders", [("F", "y", MILK)] * 3)])
        with mock.patch.object(export, "ITEM_CAP", 2):
            wb = self.xlsx(conn)
        for sheet in ("Meal templates", "Library"):
            rows = list(wb[sheet].values)
            self.assertEqual(len(rows), 4, sheet)   # header, 2 kept, the note
            self.assertIn("first 2", rows[3][0], sheet)

    def test_an_item_with_an_empty_name_keeps_its_numbers(self):
        # EX2-6
        wb = self.xlsx(FakeConn([("FROM folders", [("Dairy", "", MILK)])]))
        self.assertEqual(list(wb["Library"].values)[1][3], 125)

    def test_templates_with_the_same_name_stay_apart(self):
        # EX-9
        conn = FakeConn([])
        self.get(conn, "csv")
        sql = [s for s, _ in conn.executed if "FROM meal_templates" in s][0]
        self.assertIn("ORDER BY t.name, t.template_id, i.created_at", sql)

    def test_repeated_exports_are_rate_limited(self):
        # EX-2: every export past the limit is refused, not only the one that crosses it
        codes = [self.get(FakeConn([]), "csv").status_code for _ in range(8)]
        self.assertEqual(codes, [200] * 5 + [429] * 3)

    def test_a_busy_server_answers_503_instead_of_waiting_forever(self):
        # EX2-2: waiting exports must not park the shared worker threads
        with mock.patch.object(export, "_building", threading.BoundedSemaphore(1)), \
             mock.patch.object(export, "BUILD_WAIT", 0.05):
            export._building.acquire()
            r = self.get(FakeConn([]), "csv")
        self.assertEqual((r.status_code, r.json()["detail"]["error_type"]), (503, "busy"))
        self.assertIn("retry-after", r.headers)

    def test_a_failed_export_does_not_use_up_a_slot(self):
        # EX2-5: a cold database must not lock the user out of the one thing the 15 days are for
        def boom(sql, params):
            raise RuntimeError("neon waking up")
        client = route(self, FakeConn([("FROM daily_log", boom)]))
        for _ in range(6):
            client.get("/export?format=csv", headers={"Authorization": "Bearer x"})
        self.assertEqual(self.get(FakeConn([]), "csv").status_code, 200)

    def test_the_limit_holds_under_a_parallel_burst(self):
        # EX2-4: check and record are one step
        results = []
        def go():
            results.append(export._allowed("user-9") is not None)
        threads = [threading.Thread(target=go) for _ in range(40)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(results.count(True), export.PER_WINDOW)

    def test_the_browser_may_read_the_filename(self):
        # EX-3: the app downloads with fetch (it must send the token), so the header has to be exposed
        r = self.get(FakeConn([]), "csv", Origin="https://nutritional-tracker-delta.vercel.app")
        self.assertIn("content-disposition", r.headers.get("access-control-expose-headers", "").lower())

    def test_export_never_imports_main_itself(self):
        # EX-4: main sets export.m, like api_v1.m; an import would load a second main under `python main.py`
        with open(os.path.join(os.path.dirname(export.__file__), "export.py"), encoding="utf-8") as f:
            self.assertNotIn("import main", f.read())
        self.assertIs(export.m, main)


if __name__ == "__main__":
    unittest.main()
